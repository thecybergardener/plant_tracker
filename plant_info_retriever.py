from plant import Plant, PlantCollection
from openpyxl import load_workbook
from datetime import datetime
from dateutil.relativedelta import relativedelta # To handle date difference in years, months, days

class PlantInfoRetriever(PlantCollection):
    def __init__(self, excel_file_path) -> None:
        super().__init__()
        self.excel_file_path = excel_file_path
        self.get_all_plants()
     
    def get_all_plants(self):
        wb = load_workbook(self.excel_file_path)
        sheet = wb['Basic info']

        for row in sheet.iter_rows(min_row=2, values_only=True):
            plant_name = row[1]
            species = row[0]
            purchase_date = row[2]
            description = row[3] if row[3] else ""
            plant = Plant(plant_name, species, purchase_date, description)
            self.plant_collection.append(plant)
   
    # Function to calculate years, months, and days of ownership
    def calculate_ownership(self, purchase_date) -> str:
        today = datetime.now().date()
        if purchase_date:
            ownership_duration = relativedelta(today, purchase_date)
            years = ownership_duration.years
            months = ownership_duration.months
            days = ownership_duration.days
            # Format the purchase date as MM/DD/YYYY
            formatted_purchase_date = purchase_date.strftime("%m/%d/%Y")
            return f"{years} years, {months} months, {days} days (Purchased on {formatted_purchase_date})"
        return "N/A"