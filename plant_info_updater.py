from plant import Plant
from openpyxl import load_workbook

class PlantInfoUpdater(Plant):
    def __init__(self, excel_file_path):
        super().__init__()
        self.excel_file_path = excel_file_path

    def update_plant_info(self, plant_name, date, watering, humidity, temperature, notes):
        wb = load_workbook(self.excel_file_path)
        current_month = date.strftime("%B_%Y")

        # Ensure the sheet exists for the current month
        if current_month not in wb.sheetnames:
            new_sheet = wb.create_sheet(current_month)
            new_sheet.append(["Date", "Plant Name", "Watering", "Humidity", "Temperature", "Notes"])
        
        sheet = wb[current_month]
        row_data = [date.strftime("%m/%d/%Y"), plant_name, watering, humidity, temperature, notes]
        sheet.append(row_data)

        wb.save(self.excel_file_path)

    def add_tracker_details(self):
        pass